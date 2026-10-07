from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT = OUTPUT_DIR / "Инструкция_для_тестировщиков_Duet.docx"
LOGO = ROOT / "assets" / "duet-birds.png"

INK = "34272E"
MUTED = "71656A"
PRIMARY = "9D536D"
SECONDARY = "72465E"
PALE = "FBF3F5"
PALE_ALT = "F7F7F7"
BORDER = "D9D9D9"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, *, bold=False, color=INK, size=None, italic=False):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    run.italic = italic
    return run


def add_bullet(doc, text, *, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    add_text(paragraph, text, size=10.5)
    return paragraph


def add_numbered(doc, text, number):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.65)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    add_text(paragraph, f"{number}.  ", size=10.5)
    add_text(paragraph, text, size=10.5)
    return paragraph


def add_heading(doc, text, level=1, *, force_new_page=False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = force_new_page
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, size=16 if level == 1 else 12.5, bold=True, color="000000")
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.13
    if bold_lead and text.startswith(bold_lead):
        add_text(paragraph, bold_lead, bold=True, size=10.5)
        add_text(paragraph, text[len(bold_lead):], size=10.5)
    else:
        add_text(paragraph, text, size=10.5)
    return paragraph


def add_task(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.25)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    add_text(paragraph, "□  ", bold=True, color=PRIMARY, size=12)
    add_text(paragraph, text, size=10.5)
    return paragraph


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(6)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(INK)


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Duet  ·  инструкция для тестирования")
    set_font(run, size=8.5, color=MUTED)


def add_schedule(doc):
    rows = [
        ("День 1", "Подключиться к тесту, установить Duet и создать аккаунт", "5–10 минут"),
        ("Дни 2–4", "Попробовать пространство пары и добавить важную дату", "5–10 минут"),
        ("Дни 5–8", "Добавить желание, заметку или чек-лист", "5–10 минут"),
        ("Дни 9–12", "Проверить вопрос дня, статистику или напоминание", "5–10 минут"),
        ("Дни 13–14", "Открыть приложение ещё раз и прислать короткий итог", "5 минут"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(2.8), Cm(10.4), Cm(2.8)]
    for cell, text, width in zip(table.rows[0].cells, ("Когда", "Что сделать", "Время"), widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, SECONDARY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(paragraph, text, bold=True, color="FFFFFF", size=9.5)
    set_repeat_table_header(table.rows[0])

    for index, row_data in enumerate(rows):
        row = table.add_row()
        for cell, text, width in zip(row.cells, row_data, widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, PALE if index % 2 == 0 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if cell is not row.cells[1] else WD_ALIGN_PARAGRAPH.LEFT
            add_text(paragraph, text, bold=cell is row.cells[0], size=9.5)
    table.rows[0].cells[0].width = widths[0]
    table.rows[0].cells[1].width = widths[1]
    table.rows[0].cells[2].width = widths[2]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_footer(doc.sections[0])

    logo_paragraph = doc.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_after = Pt(4)
    picture = logo_paragraph.add_run().add_picture(str(LOGO), width=Cm(4.3))
    inline = picture._inline
    doc_pr = inline.docPr
    doc_pr.set("name", "Логотип Duet")
    doc_pr.set("descr", "Две птицы на фоне закатного солнца")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Тестирование Duet")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    add_text(subtitle, "Простая инструкция на 14 дней", bold=True, color=PRIMARY, size=13)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Cm(1.2)
    intro.paragraph_format.right_indent = Cm(1.2)
    intro.paragraph_format.space_after = Pt(13)
    intro.paragraph_format.line_spacing = 1.18
    add_text(
        intro,
        "Спасибо, что согласились помочь мне проверить Duet перед публикацией в Google Play. "
        "Пользуйтесь приложением как обычно и пишите, если что-то непонятно, неудобно или не работает.",
        size=11,
    )

    key = doc.add_paragraph()
    key.alignment = WD_ALIGN_PARAGRAPH.CENTER
    key.paragraph_format.left_indent = Cm(1.1)
    key.paragraph_format.right_indent = Cm(1.1)
    key.paragraph_format.space_after = Pt(18)
    key.paragraph_format.line_spacing = 1.15
    add_text(key, "Самое важное  ", bold=True, color=PRIMARY, size=11)
    add_text(
        key,
        "оставайтесь участником закрытого теста все 14 дней. Каждый день проверять приложение не нужно.",
        bold=True,
        size=11,
    )

    add_heading(doc, "Как подключиться")
    for number, text in enumerate((
        "Откройте ссылку на тестирование, которую я отправил вместе с этим документом.",
        "Убедитесь, что в браузере выбран тот же аккаунт Google, который используется в Google Play на телефоне.",
        "Нажмите кнопку участия в тестировании и затем откройте ссылку на приложение в Google Play.",
        "Установите Duet и создайте аккаунт. Для регистрации понадобится доступ к указанному email.",
        "До окончания теста не нажимайте выход из программы тестирования. По возможности не удаляйте приложение.",
    ), 1):
        add_numbered(doc, text, number)

    add_body(
        doc,
        "Если Google Play пишет, что приложение недоступно, пришлите мне скриншот и адрес аккаунта Google, который вы используете в магазине. Пароль присылать не нужно.",
    )

    add_heading(doc, "Как будет проходить тест")
    add_body(
        doc,
        "Достаточно открыть приложение несколько раз в течение двух недель. Дни ниже можно немного менять: главное, не выходить из теста и попробовать несколько разных функций.",
    )
    add_schedule(doc)

    add_heading(doc, "Первая короткая проверка", force_new_page=True)
    add_body(doc, "Начните с регистрации и первого знакомства с приложением.")
    for text in (
        "Создайте аккаунт, подтвердите email, если придёт письмо, и войдите.",
        "Осмотритесь на главном экране без подсказок. Найдите разделы с датами, желаниями и другими функциями.",
        "Проверьте, всё ли помещается на экране, удобно ли читать текст и понятно ли, куда нажимать.",
    ):
        add_task(doc, text)

    add_heading(doc, "Пространство пары")
    add_body(
        doc,
        "Этот сценарий лучше проверить вдвоём. Если второго участника нет, пропустите раздел и продолжайте пользоваться личным пространством.",
    )
    for text in (
        "Первый участник нажимает «Пригласить партнёра», создаёт код и отправляет его второму.",
        "Второй участник вводит шестизначный код и подтверждает объединение пространств.",
        "Проверьте, что на главной появились имена обоих участников.",
        "Добавьте или измените данные на одном телефоне и посмотрите, появились ли изменения на втором.",
    ):
        add_task(doc, text)
    add_body(
        doc,
        "Код действует 24 часа. Передавайте его только выбранному партнёру: после подключения содержимое пространств станет общим.",
        bold_lead="Код действует 24 часа.",
    )

    add_heading(doc, "Важные даты")
    for text in (
        "Откройте раздел «Даты» и добавьте событие, например поездку или годовщину.",
        "Выберите дату, категорию, значок и повторение. Сохраните событие.",
        "Откройте его снова, измените название или дату и сохраните изменения.",
        "Для прошедшего события по желанию нажмите «Состоялось». Это влияет на статистику.",
    ):
        add_task(doc, text)

    add_heading(doc, "Желания")
    for text in (
        "Откройте «Ещё», затем «Желания», или нажмите на конверт на главной.",
        "Добавьте желание в один из списков.",
        "Откройте его и по желанию добавьте описание, фотографию или комментарий.",
        "Отметьте желание исполненным и проверьте, изменились ли данные в статистике.",
    ):
        add_task(doc, text)

    add_heading(doc, "Заметки и списки дел", force_new_page=True)
    for text in (
        "Откройте «Ещё», затем «Заметки».",
        "Создайте заметку или чек-лист из нескольких пунктов и сохраните его.",
        "Закройте заметку, откройте снова и попробуйте поиск по слову из текста.",
        "По желанию переместите заметку в корзину и восстановите её.",
    ):
        add_task(doc, text)

    add_heading(doc, "Вопрос дня")
    add_body(doc, "Для полной проверки нужны два подключённых участника.")
    for text in (
        "Оба участника открывают «Ещё», затем «Вопрос дня», и включают ежедневные вопросы.",
        "Каждый пишет свой ответ. После первого ответа чужой текст ещё не должен открываться.",
        "После второго ответа должны появиться оба ответа.",
        "Откройте историю и убедитесь, что пройденный вопрос сохранился.",
    ):
        add_task(doc, text)

    add_heading(doc, "Что ещё можно попробовать")
    for text in (
        "Добавить или изменить дату начала отношений и проверить счётчик дней на главной.",
        "Для будущей даты включить напоминание на ближайшее удобное время.",
        "Открыть статистику и сравнить месяц, год и всё время.",
        "Полностью закрыть приложение, открыть его снова и проверить, сохранились ли вход и данные.",
        "Ненадолго отключить интернет и посмотреть, понятно ли приложение сообщает о проблеме.",
    ):
        add_bullet(doc, text)

    add_body(
        doc,
        "Необязательно выполнять все пункты. Если устали или столкнулись с ошибкой, зафиксируйте её и остановитесь.",
        bold_lead="Необязательно выполнять все пункты.",
    )

    add_heading(doc, "Что важно заметить")
    for text in (
        "непонятно, куда нажать или что означает текст;",
        "нужную функцию трудно найти;",
        "кнопка не реагирует или приходится нажимать несколько раз;",
        "данные пропали, задвоились или отличаются на двух телефонах;",
        "текст обрезан, экран выглядит неровно или клавиатура закрывает поле;",
        "приложение зависло, закрылось само или показало техническое сообщение;",
        "уведомление не пришло, пришло не вовремя или открыло не тот экран.",
    ):
        add_bullet(doc, text)

    add_heading(doc, "Как прислать отзыв", force_new_page=True)
    add_body(
        doc,
        "Пишите мне удобным способом. Не нужно составлять официальный отчёт. Достаточно описать, где вы были, что сделали и что произошло.",
    )

    example = doc.add_table(rows=7, cols=2)
    example.alignment = WD_TABLE_ALIGNMENT.CENTER
    example.autofit = False
    entries = (
        ("Где", "Желания, созданное желание"),
        ("Что сделал", "Добавил фотографию и нажал «Сохранить»"),
        ("Что произошло", "Загрузка не закончилась"),
        ("Что ожидал", "Желание должно было сохраниться"),
        ("Телефон", "Samsung S23, Android 16"),
        ("Повторилось", "Да, два раза"),
    )
    widths = [Cm(4), Cm(12)]
    header_row = example.rows[0]
    for cell, text, width in zip(header_row.cells, ("Что указать", "Пример"), widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, SECONDARY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, text, bold=True, color="FFFFFF", size=9.8)
    set_repeat_table_header(header_row)
    for row_index, (label, value) in enumerate(entries, 1):
        row = example.rows[row_index]
        for cell, text, width in zip(row.cells, (label, value), widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, PALE if cell is row.cells[0] else (PALE_ALT if row_index % 2 else "FFFFFF"))
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_text(paragraph, text, bold=cell is row.cells[0], size=9.8)

    add_body(
        doc,
        "Если возможно, приложите скриншот или короткую запись экрана. При неожиданном закрытии приложения напишите, что делали прямо перед этим.",
    )

    add_heading(doc, "Четыре вопроса в конце")
    for number, text in enumerate((
        "Что в приложении понравилось больше всего?",
        "Что было труднее всего понять или найти?",
        "Какая функция показалась лишней или какой функции не хватило?",
        "Стали бы вы пользоваться Duet дальше и почему?",
    ), 1):
        add_numbered(doc, text, number)

    add_heading(doc, "Безопасность данных")
    for text in (
        "Не присылайте мне пароль, код подтверждения из письма или другие секретные данные.",
        "Для теста можно использовать вымышленные даты, желания, фотографии и заметки.",
        "Не удаляйте аккаунт, если я отдельно не попросил проверить удаление. Это действие необратимо.",
        "Код пары отправляйте только человеку, с которым хотите объединить пространство.",
    ):
        add_bullet(doc, text)

    thanks = doc.add_paragraph()
    thanks.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thanks.paragraph_format.space_before = Pt(16)
    thanks.paragraph_format.space_after = Pt(0)
    add_text(thanks, "Спасибо за помощь с Duet", bold=True, color=PRIMARY, size=13)

    doc.core_properties.title = "Тестирование Duet"
    doc.core_properties.subject = "Инструкция для участников закрытого тестирования Duet"
    doc.core_properties.author = "Duet"
    doc.core_properties.keywords = "Duet, тестирование, Google Play"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
